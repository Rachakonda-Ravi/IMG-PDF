import os
import re
from pypdf import PdfReader
from docx import Document

def clean_text(text):
    """
    Remove NULL bytes and ASCII control characters that XML (python-docx) cannot handle.
    Keeps standard whitespace like newlines (\n), carriage returns (\r), and tabs (\t).
    """
    if not text:
        return ""
    # Matches control characters excluding \t (0x09), \n (0x0A), \r (0x0D)
    return re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x84\x86-\x9F]', '', text)

def convert_pdf_to_doc():
    # Get the directory where this script is located
    script_dir = os.path.dirname(os.path.abspath(__file__))
    
    # Prompt user for the PDF file path
    pdf_path = input("Enter the path to the PDF file: ").strip(' "\'')
    
    if not os.path.exists(pdf_path):
        print(f"Error: File not found at '{pdf_path}'")
        return

    # Set output filename based on the PDF name
    pdf_name = os.path.splitext(os.path.basename(pdf_path))[0]
    output_docx_path = os.path.join(script_dir, f"{pdf_name}.docx")

    print("\nReading PDF and converting...")
    reader = PdfReader(pdf_path)
    total_pages = len(reader.pages)
    print(f"Total pages detected: {total_pages}\n")

    doc = Document()
    doc.add_heading(clean_text(pdf_name), level=0)

    for i, page in enumerate(reader.pages, start=1):
        raw_text = page.extract_text() or ""
        cleaned_text = clean_text(raw_text)
        
        # Add heading for page indicator
        doc.add_heading(f"--- Page {i} ---", level=2)
        
        if cleaned_text.strip():
            doc.add_paragraph(cleaned_text)
        else:
            doc.add_paragraph("[No readable text found on this page]")
            
        # Insert a page break between PDF pages
        if i < total_pages:
            doc.add_page_break()
            
        # Dynamically updates on the same line in terminal/command prompt
        print(f"\rProcessing page {i}/{total_pages}...", end="", flush=True)

    doc.save(output_docx_path)
    print(f"\n\nDone! Saved output to: {output_docx_path}")

if __name__ == "__main__":
    convert_pdf_to_doc()